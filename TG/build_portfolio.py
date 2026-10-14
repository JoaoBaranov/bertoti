from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(3)
sec.left_margin = Cm(3)
sec.bottom_margin = Cm(2)
sec.right_margin = Cm(2)
for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[style_name]
    style.font.name = 'Times New Roman'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
doc.styles['Normal'].font.size = Pt(12)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 2'].font.size = Pt(12)

def para(text='', center=False):
    x = doc.add_paragraph(text)
    x.paragraph_format.line_spacing = 1.5
    x.paragraph_format.first_line_indent = Cm(1.25)
    if center:
        x.alignment = WD_ALIGN_PARAGRAPH.CENTER

def h(text, level=1):
    doc.add_heading(text, level=level)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('JOÃO VITOR ARAUJO BARANOV')
for text in ['', 'PORTFÓLIO DE APRENDIZAGEM DE PROJETOS INTEGRADOS', 'PROJETOS API 1 A 4', '', 'São José dos Campos', '2026']:
    para(text, True)
doc.add_page_break()

h('Resumo')
para('Este portfólio reúne quatro projetos integradores desenvolvidos ao longo da graduação: uma calculadora científica, um gerenciador de horários, uma plataforma de competências e um sistema de gerenciamento de manutenções. O documento apresenta os problemas, as soluções, as tecnologias e a participação do autor nas etapas de requisitos, organização, implementação, testes e documentação.')
para('Palavras-chave: projetos integradores; desenvolvimento de software; metodologias ágeis; portfólio acadêmico.')
h('Sumário')
for item in ['1 Introdução', '2 Projetos', '2.1 API-Concord', '2.2 Concord Grid Manager', '2.3 Proficio', '2.4 Trivio', '3 Considerações finais', 'Referências']:
    para(item)
doc.add_page_break()

h('1 Introdução')
para('Os projetos integradores aproximam os conteúdos da graduação de situações práticas de desenvolvimento. Em cada semestre, a equipe precisou compreender um problema, organizar as atividades, definir uma solução e entregar uma aplicação ou protótipo funcional. A sequência dos projetos mostra a evolução da experiência com lógica de programação, desenvolvimento web, bancos de dados, documentação e práticas ágeis.')
para('Este documento apresenta os quatro primeiros projetos selecionados para o portfólio. Em todos eles, a participação ocorreu de forma transversal, com colaboração em diferentes fases do trabalho e adaptação às necessidades da equipe.')

projetos = [
('2.1 API-Concord', 'O API-Concord criou uma calculadora científica em Portugol, com menu para seleção das operações. A aplicação reúne adição, subtração, multiplicação, divisão, fatorial, função de segundo grau, concatenação de strings, juros simples e compostos e conversões entre bases numéricas. Participei do levantamento das necessidades, organização das atividades, implementação, testes e documentação. O projeto fortaleceu minha base de lógica de programação e validação de resultados. Tecnologia: Portugol. Repositório: https://github.com/RichardCordeiro/API-Concord.'),
('2.2 Concord Grid Manager', 'O Concord Grid Manager organiza horários de aula sem sobreposição, considerando disciplinas, professores, cursos, turmas e a carga semanal. O sistema permite cadastros, montagem da grade, edição manual, validação de conflitos e organização do semestre. Participei do entendimento dos requisitos, backlog, implementação, testes, documentação e revisão das entregas. Tecnologias: JavaFX, Java, Maven, Scene Builder e MySQL. Repositório: https://github.com/Concord-API/API-2.'),
('2.3 Proficio', 'O Proficio é uma plataforma de mapeamento de competências, concebida como um LinkedIn interno. Colaboradores mantêm perfis com competências, soft skills, certificações e experiências; gestores podem complementar perfis, buscar habilidades e analisar competências. Participei do levantamento de requisitos, backlog, desenvolvimento, testes, documentação e revisão em três sprints. Tecnologias: Java, Spring, TypeScript, JavaScript, HTML, CSS, MySQL, Docker, Git e GitHub. Repositório: https://github.com/Concord-API/API-3.'),
('2.4 Trivio', 'O Trivio gerencia manutenções em equipamentos distribuídos globalmente. A solução contempla técnicos, equipamentos, requisitos, prioridades, localizações, manutenções recorrentes, documentação, alertas e mapa. Participei do entendimento do problema, backlog, implementação, testes, documentação e revisão das entregas. Tecnologias: Java, Spring, JavaScript, TypeScript, HTML, CSS, Vue.js, Oracle Database, Docker, Git e GitHub. Repositório: https://github.com/Concord-API/API-4.')]
for title_text, body in projetos:
    h(title_text, 2)
    para(body)
    if title_text.startswith('2.1'):
        para('Contribuições pessoais: participei do entendimento do problema, organização das atividades, implementação das operações, testes e documentação.')
        para('Hard skills: lógica de programação, Portugol, estruturas condicionais, funções e GitHub.')
        para('Soft skills: trabalho em equipe, comunicação, organização, resolução de problemas e adaptação.')
    elif title_text.startswith('2.2'):
        para('Contribuições pessoais: participei do entendimento dos requisitos, backlog, implementação, testes, documentação e revisão das entregas.')
        para('Hard skills: Java, JavaFX, MySQL, modelagem de dados, regras de negócio, Maven e Git.')
        para('Soft skills: trabalho colaborativo, comunicação, organização de sprint, análise de conflitos e documentação técnica.')
    elif title_text.startswith('2.3'):
        para('Contribuições pessoais: participei do levantamento de requisitos, priorização do backlog, desenvolvimento, testes, documentação e revisão das entregas.')
        para('Hard skills: Java, Spring, TypeScript, JavaScript, HTML, CSS, MySQL, Docker e desenvolvimento de APIs.')
        para('Soft skills: comunicação, organização de backlog, trabalho em sprints, responsabilidade compartilhada e adaptação.')
    else:
        para('Contribuições pessoais: participei do entendimento do problema, backlog, implementação, testes, documentação e revisão das entregas.')
        para('Hard skills: Java, Spring, Vue.js, JavaScript, TypeScript, Oracle Database, Docker e integração entre camadas.')
        para('Soft skills: trabalho em equipe, planejamento, comunicação técnica, organização de prioridades e resolução de problemas.')

h('3 Considerações finais')
para('Os quatro projetos mostram uma evolução gradual da minha formação. O primeiro consolidou fundamentos de lógica; os projetos seguintes ampliaram a experiência com regras de negócio, interfaces, bancos de dados, desenvolvimento web, documentação e organização de equipes em sprints.')
para('A participação em várias etapas desenvolveu comunicação, colaboração, adaptação e resolução de problemas. Trabalhar com demandas diferentes ajudou a compreender que uma boa solução depende da implementação, do entendimento do usuário, da organização do trabalho e da validação das entregas.')
h('Referências')
for ref in ['CORDEIRO, Richard et al. API-Concord. GitHub. Disponível em: https://github.com/RichardCordeiro/API-Concord. Acesso em: 14 out. 2026.', 'CONCORD-API. API-2. GitHub. Disponível em: https://github.com/Concord-API/API-2. Acesso em: 14 out. 2026.', 'CONCORD-API. API-3. GitHub. Disponível em: https://github.com/Concord-API/API-3. Acesso em: 14 out. 2026.', 'CONCORD-API. API-4. GitHub. Disponível em: https://github.com/Concord-API/API-4. Acesso em: 14 out. 2026.', 'FATEC PROFESSOR JESSEN VIDAL. Modelo de portfólio de aprendizagem de projetos integrados. São José dos Campos, 2026.']:
    para(ref)
doc.save('portfolio-academico.docx')
